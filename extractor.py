"""
extractor.py — runs INSIDE the browser via Pyodide (Python compiled to
WebAssembly). This is the same tested extraction/breakdown/pivot logic as
the desktop app, with the desktop-only bits removed (Tkinter GUI, the
optional spaCy NLP fallback, and server-side MSAL auth — cloud sign-in is
handled by JavaScript instead, since browser sign-in is simpler there).

Nothing in this file makes a network call. It only ever touches files that
JavaScript has already written into Pyodide's in-memory filesystem.
"""

import os
import re
import json
import sqlite3
from datetime import datetime, date

import docx  # provided via micropip in the browser (python-docx + lxml)

DB_PATH = "/data/indicator_data.db"

DEFAULT_INDICATORS = [
    {"name": "Completion Rate", "synonyms": ["Completion %", "Completion Percentage"]},
    {"name": "Budget Utilized", "synonyms": ["Budget Used", "Budget Spent"]},
]


# --------------------------------------------------------------------------
# Matching helpers (identical to the desktop app)
# --------------------------------------------------------------------------

def normalize(text):
    text = text.strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def all_labels_for(indicator):
    labels = [indicator["name"]] + indicator.get("synonyms", [])
    return [normalize(l) for l in labels if l.strip()]


def find_label_match(cell_text, indicators):
    norm = normalize(cell_text)
    norm_stripped = re.sub(r"[:\-–]\s*$", "", norm).strip()
    for ind in indicators:
        for label in all_labels_for(ind):
            if norm_stripped == label:
                return ind
    return None


def extract_inline_value(text, indicators):
    norm = normalize(text)
    for ind in indicators:
        for label in all_labels_for(ind):
            pattern = r"^" + re.escape(label) + r"\s*[:\-–]\s*(.+)$"
            m = re.match(pattern, norm)
            if m:
                sep_match = re.search(r"[:\-–]", text)
                if sep_match:
                    value = text[sep_match.end():].strip()
                    if value:
                        return ind, value
    return None, None


# --------------------------------------------------------------------------
# Breakdown detection: population type / gender / unit (identical logic)
# --------------------------------------------------------------------------

BREAKDOWN_KEYWORDS = [
    ("host community", ("population_type", "Host")),
    ("refugees", ("population_type", "Refugee")),
    ("refugee", ("population_type", "Refugee")),
    ("host", ("population_type", "Host")),
    ("idps", ("population_type", "IDP")),
    ("idp", ("population_type", "IDP")),
    ("returnees", ("population_type", "Returnee")),
    ("returnee", ("population_type", "Returnee")),
    ("asylum seekers", ("population_type", "Asylum Seeker")),
    ("asylum seeker", ("population_type", "Asylum Seeker")),
    ("females", ("gender", "Female")),
    ("female", ("gender", "Female")),
    ("women", ("gender", "Female")),
    ("girls", ("gender", "Female")),
    ("males", ("gender", "Male")),
    ("male", ("gender", "Male")),
    ("men", ("gender", "Male")),
    ("boys", ("gender", "Male")),
    ("households", ("unit", "Households")),
    ("household", ("unit", "Households")),
    ("hh", ("unit", "Households")),
    ("individuals", ("unit", "Individuals")),
    ("individual", ("unit", "Individuals")),
    ("persons", ("unit", "Individuals")),
    ("person", ("unit", "Individuals")),
    ("people", ("unit", "Individuals")),
]

_NUMBER_PATTERN = re.compile(r"\$?\d{1,3}(?:,\d{3})+(?:\.\d+)?%?|\$?\d+(?:\.\d+)?%?")
_TOTAL_RESET_KEYWORDS = ["total", "overall", "combined", "grand total", "aggregate", "all groups"]


def classify_breakdown_text(text):
    lower = " " + normalize(text) + " "
    found = {"population_type": None, "gender": None, "unit": None}
    best_end = {"population_type": -1, "gender": -1, "unit": -1}

    for keyword, (category, label) in BREAKDOWN_KEYWORDS:
        pattern = r"(?<![a-z])" + re.escape(keyword) + r"(?![a-z])"
        last_match = None
        for m in re.finditer(pattern, lower):
            last_match = m
        if last_match and last_match.end() > best_end[category]:
            best_end[category] = last_match.end()
            found[category] = label

    total_end = -1
    for kw in _TOTAL_RESET_KEYWORDS:
        pattern = r"(?<![a-z])" + re.escape(kw) + r"(?![a-z])"
        last_match = None
        for m in re.finditer(pattern, lower):
            last_match = m
        if last_match:
            total_end = max(total_end, last_match.end())

    if total_end != -1:
        for category in found:
            if total_end > best_end[category]:
                found[category] = None

    return found


def extract_breakdown_values(text):
    rows = []
    for num_match in _NUMBER_PATTERN.finditer(text):
        num_start = num_match.start()
        window_start = max(0, num_start - 35)
        preceding = text[window_start:num_start]
        tags = classify_breakdown_text(preceding)
        rows.append({
            "population_type": tags["population_type"],
            "gender": tags["gender"],
            "unit": tags["unit"],
            "value": num_match.group(0).strip(),
        })
    return rows


def merge_tags(*dicts):
    merged = {"population_type": None, "gender": None, "unit": None}
    for d in dicts:
        if not d:
            continue
        for k in merged:
            if merged[k] is None and d.get(k):
                merged[k] = d[k]
    return merged


def make_result(filename, indicator_name, value, source, tags=None):
    tags = tags or {}
    return {
        "document": filename, "indicator": indicator_name, "value": value, "source": source,
        "population_type": tags.get("population_type") or "",
        "gender": tags.get("gender") or "",
        "unit": tags.get("unit") or "",
    }


def _has_digit(text):
    return bool(re.search(r"\d", text))


# --------------------------------------------------------------------------
# Extraction core (identical logic; use_nlp path removed — no spaCy in-browser)
# --------------------------------------------------------------------------

def extract_from_lines(lines, indicators, filename, source_label):
    """
    Shared matching pass over a flat list of text lines — used for both
    document paragraphs and OCR output (split into lines), so a value found
    via OCR gets identical matching/breakdown-tagging as one found in a
    native paragraph. No NLP fallback in the browser build (see README).
    """
    results = []
    for i, text in enumerate(lines):
        if not text.strip():
            continue

        ind, value = extract_inline_value(text, indicators)
        if ind:
            line_tags = classify_breakdown_text(text)
            tagged_values = extract_breakdown_values(value)
            if tagged_values:
                for tv in tagged_values:
                    merged = merge_tags(tv, line_tags)
                    results.append(make_result(filename, ind["name"], tv["value"], source_label, merged))
            else:
                results.append(make_result(filename, ind["name"], value, source_label, line_tags))
            continue

        matched = find_label_match(text, indicators)
        if matched:
            label_tags = classify_breakdown_text(text)
            for j in range(i + 1, min(i + 3, len(lines))):
                nxt = lines[j].strip()
                if nxt:
                    tagged_values = extract_breakdown_values(nxt)
                    if tagged_values:
                        for tv in tagged_values:
                            merged = merge_tags(tv, label_tags)
                            results.append(make_result(filename, matched["name"], tv["value"], source_label, merged))
                    else:
                        merged = merge_tags(classify_breakdown_text(nxt), label_tags)
                        results.append(make_result(filename, matched["name"], nxt, source_label, merged))
                    break
            continue

    return results


def extract_from_docx(path, indicators):
    """
    path: a file path already written into Pyodide's filesystem by JS.
    indicators: list of {"name": ..., "synonyms": [...]}.
    Returns list of dicts: document, indicator, value, source,
    population_type, gender, unit.
    """
    results = []
    doc = docx.Document(path)
    filename = path.rsplit("/", 1)[-1]

    paragraphs = [p.text for p in doc.paragraphs]
    results.extend(extract_from_lines(paragraphs, indicators, filename, "paragraph"))

    for table in doc.tables:
        grid = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        n_rows = len(grid)
        header_row = grid[0] if n_rows else []

        for r, row in enumerate(grid):
            for c, cell_text in enumerate(row):
                if not cell_text:
                    continue

                ind, value = extract_inline_value(cell_text, indicators)
                if ind:
                    tags = classify_breakdown_text(cell_text)
                    results.append(make_result(filename, ind["name"], value, "table", tags))
                    continue

                matched = find_label_match(cell_text, indicators)
                if not matched:
                    continue

                value_cols = [
                    c2 for c2 in range(len(row))
                    if c2 != c and row[c2].strip() and _has_digit(row[c2])
                ]

                if value_cols:
                    for c2 in value_cols:
                        value_text = row[c2].strip()
                        header_tags = classify_breakdown_text(header_row[c2]) if c2 < len(header_row) else {}
                        exclude = {c, c2}
                        row_context = " ".join(
                            row[i] for i in range(len(row)) if i not in exclude and row[i].strip()
                        )
                        row_tags = classify_breakdown_text(row_context)
                        value_tags = classify_breakdown_text(value_text)
                        merged = merge_tags(value_tags, row_tags, header_tags)
                        results.append(make_result(filename, matched["name"], value_text, "table", merged))
                    continue

                if (r + 1 < n_rows and c < len(grid[r + 1]) and grid[r + 1][c].strip()
                        and _has_digit(grid[r + 1][c])):
                    value_text = grid[r + 1][c].strip()
                    header_tags = classify_breakdown_text(header_row[c]) if c < len(header_row) else {}
                    next_row = grid[r + 1]
                    next_row_context = " ".join(
                        next_row[i] for i in range(len(next_row)) if i != c and next_row[i].strip()
                    )
                    label_row_context = " ".join(
                        row[i] for i in range(len(row)) if i != c and row[i].strip()
                    )
                    row_tags = classify_breakdown_text(next_row_context + " " + label_row_context)
                    value_tags = classify_breakdown_text(value_text)
                    merged = merge_tags(value_tags, row_tags, header_tags)
                    results.append(make_result(filename, matched["name"], value_text, "table", merged))

    return results


# --------------------------------------------------------------------------
# Reporting period helpers
# --------------------------------------------------------------------------

PERIOD_TYPES = ["All", "Year", "Month", "Week"]


def suggested_period_value(period_type):
    today = date.today()
    if period_type == "Year":
        return str(today.year)
    if period_type == "Month":
        return today.strftime("%Y-%m")
    if period_type == "Week":
        return today.strftime("%Y-W%V")
    return ""


def period_label(period_type, period_value):
    if not period_type or period_type == "All":
        return "All"
    return f"{period_value} ({period_type})" if period_value else f"({period_type})"


# --------------------------------------------------------------------------
# Local persistent dataset (SQLite, backed by a file JS saves to IndexedDB)
# --------------------------------------------------------------------------

def get_db_conn():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = get_db_conn()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS extracted_data (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            document TEXT,
            indicator TEXT,
            value TEXT,
            population_type TEXT,
            gender TEXT,
            unit TEXT,
            source TEXT,
            period_type TEXT,
            period_value TEXT,
            extracted_at TEXT,
            pushed_at TEXT
        )
    """)
    conn.commit()
    conn.close()


def insert_rows(rows_json):
    """rows_json: JSON string (list of dicts) — Pyodide/JS interop is
    cleanest passing plain strings across the boundary. Appends only;
    never deletes or overwrites existing rows."""
    rows = json.loads(rows_json)
    conn = get_db_conn()
    now = datetime.now().isoformat(timespec="seconds")
    conn.executemany("""
        INSERT INTO extracted_data
        (document, indicator, value, population_type, gender, unit, source,
         period_type, period_value, extracted_at, pushed_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL)
    """, [
        (r.get("document", ""), r.get("indicator", ""), r.get("value", ""),
         r.get("population_type", ""), r.get("gender", ""), r.get("unit", ""),
         r.get("source", ""), r.get("period_type", ""), r.get("period_value", ""),
         now)
        for r in rows
    ])
    conn.commit()
    conn.close()


def fetch_all_rows_json(filters_json="{}"):
    filters = json.loads(filters_json) if filters_json else {}
    conn = get_db_conn()
    query = "SELECT * FROM extracted_data"
    params = []
    clauses = []
    for k, v in filters.items():
        if v and v != "(All)":
            clauses.append(f"{k} = ?")
            params.append(v)
    if clauses:
        query += " WHERE " + " AND ".join(clauses)
    query += " ORDER BY id DESC"
    cur = conn.execute(query, params)
    rows = [dict(row) for row in cur.fetchall()]
    conn.close()
    return json.dumps(rows)


def fetch_unpushed_rows_json():
    conn = get_db_conn()
    cur = conn.execute("SELECT * FROM extracted_data WHERE pushed_at IS NULL ORDER BY id ASC")
    rows = [dict(row) for row in cur.fetchall()]
    conn.close()
    return json.dumps(rows)


def mark_pushed(ids_json):
    ids = json.loads(ids_json)
    if not ids:
        return
    conn = get_db_conn()
    now = datetime.now().isoformat(timespec="seconds")
    conn.executemany("UPDATE extracted_data SET pushed_at=? WHERE id=?", [(now, i) for i in ids])
    conn.commit()
    conn.close()


def distinct_values_json(column):
    assert column in {
        "document", "indicator", "value", "population_type", "gender",
        "unit", "source", "period_type", "period_value",
    }, "invalid column"
    conn = get_db_conn()
    cur = conn.execute(
        f"SELECT DISTINCT {column} FROM extracted_data "
        f"WHERE {column} IS NOT NULL AND {column} != '' ORDER BY {column}"
    )
    vals = [row[0] for row in cur.fetchall()]
    conn.close()
    return json.dumps(vals)


def row_count():
    conn = get_db_conn()
    n = conn.execute("SELECT COUNT(*) FROM extracted_data").fetchone()[0]
    conn.close()
    return n


def clear_all_data():
    """Used only by the explicit 'Reset all data' button in the UI."""
    conn = get_db_conn()
    conn.execute("DELETE FROM extracted_data")
    conn.commit()
    conn.close()


# --------------------------------------------------------------------------
# Table builder (pivot-style summaries)
# --------------------------------------------------------------------------

PIVOT_FIELDS = {
    "Indicator": "indicator",
    "Document": "document",
    "Population Type": "population_type",
    "Gender": "gender",
    "Household/Individual": "unit",
    "Reporting Period": "period_value",
    "Source": "source",
}


def parse_numeric(value):
    if not value:
        return None
    cleaned = re.sub(r"[^\d.\-]", "", value)
    if cleaned in ("", "-", "."):
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def build_pivot_json(rows_json, row_field_key, col_field_key, agg_mode):
    rows = json.loads(rows_json)
    row_field = PIVOT_FIELDS[row_field_key]
    col_field = PIVOT_FIELDS.get(col_field_key) if col_field_key != "(None)" else None

    row_keys, col_keys = [], []
    grid = {}
    for r in rows:
        rk = r.get(row_field) or "(blank)"
        ck = (r.get(col_field) or "(blank)") if col_field else "Value"
        if rk not in row_keys:
            row_keys.append(rk)
        if ck not in col_keys:
            col_keys.append(ck)
        grid.setdefault(rk, {}).setdefault(ck, []).append(r.get("value", ""))

    row_keys.sort()
    col_keys.sort()

    table_rows = []
    for rk in row_keys:
        out_row = [rk]
        for ck in col_keys:
            values = grid.get(rk, {}).get(ck, [])
            if agg_mode == "Count":
                cell = str(len(values))
            elif agg_mode == "Sum":
                nums = [parse_numeric(v) for v in values]
                nums = [n for n in nums if n is not None]
                cell = f"{sum(nums):,.2f}" if nums else "-"
            else:
                cell = ", ".join(values) if values else ""
            out_row.append(cell)
        table_rows.append(out_row)

    headers = [row_field_key] + col_keys
    return json.dumps({"headers": headers, "rows": table_rows})


def extract_from_docx_json(path, indicators_json):
    indicators = json.loads(indicators_json)
    return json.dumps(extract_from_docx(path, indicators))


def save_docx_images_json(path):
    """
    Pulls every embedded picture out of a docx and writes each one to
    Pyodide's filesystem so JS can read the bytes and hand them to
    Tesseract.js for OCR (the OCR engine itself is JS/WASM, not Python —
    Python's job here is just getting the image bytes out of the docx).
    Returns a JSON list of the file paths written.
    """
    doc = docx.Document(path)
    out_paths = []
    i = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            try:
                blob = rel.target_part.blob
                ext = os.path.splitext(rel.target_part.partname)[1] or ".png"
                out_path = f"/uploads/_embedded_{i}{ext}"
                with open(out_path, "wb") as f:
                    f.write(blob)
                out_paths.append(out_path)
                i += 1
            except Exception:
                continue
    return json.dumps(out_paths)


def extract_from_lines_json(lines_json, indicators_json, filename, source_label):
    """Called from JS with OCR'd text (already split into lines) — same
    matching pass used for document paragraphs."""
    lines = json.loads(lines_json)
    indicators = json.loads(indicators_json)
    return json.dumps(extract_from_lines(lines, indicators, filename, source_label))
